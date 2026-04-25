
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
import os
import pandas as pd

#File paths
input_file=os.path.join("data", "sales_data.csv")
output_dir="output"

#Create output folder if not exists
os.makedirs(output_dir, exist_ok=True)

#--------------------------
#Step 1: Read the sales data
#---------------------------

print("Reading sales data...")
#df=pd.read_csv(input_file, sep=",", encoding="utf-8")

#Read raw file as single column first to handle messy data
df=pd.read_csv(input_file,header=None)

#Split into proper columns (single column into multiple columns)
df=df[0].str.split(",", expand=True)

#Set column names manually
df.columns=[ 
    "Order ID", "Order Date", "Region", "Sales Person",
    "Product", "Category", "Units Sold", "Unit Price",
    "Total Sales", "Payment Method"
]

#Remove Spaces
df=df.apply(lambda col: col.str.strip() if col.dtype=="object" else col)

#Remove repeated header row
df=df[df.iloc[:,0].str.lower()!="order id"]

#Reset index
df=df.reset_index(drop=True)

print("\nCleaned Data Preview:")
print(df.head())

#print("\nColumns count:", len(df.columns))
#print("\nDataset Info:")
#print(df.info())

print("\nShape:", df.shape)
print("\nMissing Values:")
print(df.isnull().sum())

print("\nDuplicte Rows:", df.duplicated().sum())

#--------------------------
#Step 2: Data Cleaning
#--------------------------

print("\nStarting Data Cleaning....")

#Treat empty strings as missing values
df=df.replace("",pd.NA)

#1. Convert numeric columns
df["Units Sold"]=pd.to_numeric(df["Units Sold"])
df["Unit Price"]=pd.to_numeric(df["Unit Price"])
df["Total Sales"]=pd.to_numeric(df["Total Sales"])

#2. Convert Order Date to proper format
df["OrderDate"]=pd.to_datetime(df["Order Date"], errors="coerce")

#3. Fill missing values
df["Sales Person"]=df["Sales Person"].fillna("Unknown")
df["Units Sold"]=df["Units Sold"].fillna(0)

#4. Remove duplicate rows
df=df.drop_duplicates()

#5. Recalculate Total Sales
df["Total Sales"]=df["Units Sold"] * df["Unit Price"]

#6. Reset index after cleaning
df=df.reset_index(drop=True)

print("\nAfter Cleaning Data Preview:")
print(df.head())

print("\nFinal Shape:", df.shape)

#-----------------------------
#Step 3: Generate Excel Report
#-----------------------------

print("\nGenerating Excel Report...")

output_file=os.path.join(output_dir,"sales_report.xlsx")

#Summary
total_sales=df["Total Sales"].sum()
total_orders=len(df)
average_order_value=df["Total Sales"].mean()

summary_df=pd.DataFrame({
    "Metric": ["Total Sales", "Total Orders", "Average Order Value"],
    "Value": [total_sales, total_orders, average_order_value]
})

#Region-wise Report
region_df=df.groupby("Region", as_index=False)["Total Sales"].sum()
region_df=region_df.sort_values(by="Total Sales", ascending=False)

#Product-wise Report
product_df=df.groupby("Product", as_index=False)["Total Sales"].sum()
product_df=product_df.sort_values(by="Total Sales", ascending=False)

#Sales Person Report
salesperson_df=df.groupby("Sales Person", as_index=False)["Total Sales"].sum()
salesperson_df=salesperson_df.sort_values(by="Total Sales", ascending=False)

#Top performer
top_sales_person=salesperson_df.iloc[0]

print("\nTop Sales Person:")
print(top_sales_person)

#Write sheets first
with pd.ExcelWriter(output_file, engine="openpyxl", mode="w") as writer:
    df.to_excel(writer, sheet_name="Cleaned Data", index=False)
    summary_df.to_excel(writer, sheet_name="Summary", index=False)
    region_df.to_excel(writer, sheet_name="Region Report", index=False)
    product_df.to_excel(writer, sheet_name="Product Report", index=False)
    salesperson_df.to_excel(writer, sheet_name="Sales Person-wise", index=False)

#Load workbook again to add charts
wb=load_workbook(output_file)

#Region chart
ws_region=wb["Region Report"]

chart1=BarChart()
chart1.title="Sales by Region"
chart1.y_axis.title="Total Sales"
chart1.x_axis.title="Region"

data1= Reference(ws_region, min_col=2, min_row=1, max_row=len(region_df)+1)
cats1= Reference(ws_region, min_col=1, min_row=2, max_row=len(region_df)+1)

chart1.add_data(data1, titles_from_data=True)
chart1.set_categories(cats1)

ws_region.add_chart(chart1, "E2")

#Product chart
ws_product=wb["Product Report"]

chart2=BarChart()
chart2.title="Sales by Product"
chart2.y_axis.title="Total Sales"
chart2.x_axis.title="Product"

data2= Reference(ws_product, min_col=2, min_row=1, max_row=len(product_df)+1)
cats2= Reference(ws_product, min_col=1, min_row=2, max_row=len(product_df)+1)

chart2.add_data(data2, titles_from_data=True)
chart2.set_categories(cats2)

ws_product.add_chart(chart2, "E2")

#Save final workbook
wb.save(output_file)

print(f"\nReport generated successfully with charts: {output_file}")


